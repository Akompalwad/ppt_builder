import zipfile

from app.schemas.presentation import CreatePresentationRequest
from app.agents.orchestrator import PresentationOrchestrator, auto_theme_for_topic, resolve_theme
from app.agents.storyline_agent import StorylineAgent
from app.agents.brief_agent import BriefInterpreterAgent
from app.agents.slide_content_agent import SlideContentAgent
from app.rendering.pptx_builder import build_presentation
from app.rendering.web_renderer import render_slide_html
from app.agents.qa_agent import PresentationQAAgent, summarize_point, topic_matches_deck
from app.schemas.presentation import LayoutType, PresentationSpec, SlideSpec
from app.agents.design_agent import DesignDirectorAgent
from app.services.image_service import ImageService
from app.services.generation_queue import GenerationQueue
from app.config import Settings
from pptx import Presentation

def test_generation_queue_snapshot_reports_waiting_depth():
    queue=GenerationQueue(max_active_jobs=1)
    with queue._condition:
        queue._active=1
        queue._pending.extend(["first:ticket", "second:ticket"])
    snapshot=queue.snapshot("second")
    assert snapshot["queue_depth"] == 3
    assert snapshot["jobs_ahead"] == 2
    assert snapshot["position"] == 2

def test_fallback_and_pptx(tmp_path):
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic="AI adoption",slide_count=4))
    assert len(spec.slides)==4
    output=build_presentation(spec,tmp_path/"deck.pptx")
    assert output.exists() and output.stat().st_size>1000

def test_qa_keeps_complete_sentence_not_ellipsis():
    result=summarize_point("Clusters require observability, security, deployment, and incident capabilities for safe operation.", 8)
    assert result.endswith(".")
    assert "…" not in result

def test_design_director_assigns_a_cover_and_close():
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic="Kubernetes versus monolithic architecture",slide_count=6))
    plan=DesignDirectorAgent().apply(spec)
    assert plan.decisions[0].layout.value == "title_slide"
    assert plan.decisions[-1].layout.value == "summary"

def test_design_director_records_a_visual_direction():
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic="A product migration roadmap",slide_count=5))
    plan=DesignDirectorAgent().apply(spec)
    assert len(plan.decisions) == 5
    assert all(decision.visual_direction for decision in plan.decisions)
    assert all(decision.background_treatment for decision in plan.decisions)

def test_qa_restores_both_sides_of_a_comparison_cover():
    spec=PresentationSpec(
        title="Compare NSDL and CDSL", topic="Compare NSDL and CDSL. Which should I choose and why?",
        slides=[SlideSpec(slide_number=1, title="NSDL vs.", subtitle="A decision guide", purpose="Set the context", layout_type=LayoutType.title_slide)],
    )
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].title == "NSDL vs. CDSL"

def test_qa_removes_unpunctuated_comparison_question_from_cover_title():
    spec=PresentationSpec(
        title="NSDL versus CDSL", topic="NSDL vs CDSL which should I choose and why",
        slides=[SlideSpec(slide_number=1, title="NSDL vs.", purpose="Choose a depository", layout_type=LayoutType.title_slide)],
    )
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].title == "NSDL vs. CDSL"

def test_qa_does_not_turn_an_overview_request_into_a_fake_comparison():
    spec=PresentationSpec(
        title="SlideWeaver", topic="Create a 7-slide product and technical overview of SlideWeaver, an AI-powered PowerPoint generator.",
        slides=[SlideSpec(slide_number=1, title="SlideWeaver", purpose="Overview", layout_type=LayoutType.title_slide)],
    )
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].title == "SlideWeaver"

def test_qa_replaces_private_story_instruction_with_slide_fact():
    spec=PresentationSpec(
        title="NSDL versus CDSL", topic="Compare NSDL and CDSL",
        slides=[SlideSpec(slide_number=1, title="Digital wealth", purpose="Set the decision context and stakes by explaining the importance of choosing the.", layout_type=LayoutType.content_with_visual, elements=[{"heading":"What is a depository?", "body":"A digital vault that holds shares and bonds."}])],
    )
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].purpose == "A digital vault that holds shares and bonds."

def test_qa_replaces_incomplete_or_planning_purpose_with_visible_fact():
    spec=PresentationSpec(
        title="Green logistics", topic="Green logistics",
        slides=[SlideSpec(
            slide_number=1, title="Delivery cost pressure", layout_type=LayoutType.summary,
            purpose="To provide clear contact information and a direct",
            elements=[{"heading":"Fuel use", "body":"Route optimization cuts fuel waste and delivery emissions."}],
        )],
    )
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].purpose == "Route optimization cuts fuel waste and delivery emissions."

def test_topic_guard_rejects_unrelated_provider_deck():
    spec=PresentationSpec(
        title="EcoRoute", topic="The History and Future of Quantum Computing",
        slides=[SlideSpec(slide_number=1, title="Green Logistics", purpose="Optimize delivery routes.", layout_type=LayoutType.title_slide)],
    )
    assert not topic_matches_deck(spec, "The History and Future of Quantum Computing")

def test_prose_constraints_preserve_timeline_concepts_and_final_count():
    prompt=("Create an educational presentation about quantum computing spanning exactly 7 slides. "
            "The content needs to move from a historical timeline (1980s to present), explain 3 complex concepts "
            "(superposition, entanglement, qubits) using nested subheadings, and finish with a slide detailing 4 distinct future industries it will disrupt.")
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=6)
    assert brief.is_structured and brief.requested_slide_count == 7
    assert brief.constraint_for_slide(2).layout_type == LayoutType.timeline
    concepts=brief.constraint_for_slide(4)
    assert concepts.requirements == ["superposition", "entanglement", "qubits"]
    assert concepts.story_stage == "concept foundations"
    assert brief.constraint_for_slide(3).story_stage == "progress to present"
    final=brief.constraint_for_slide(7)
    assert final.layout_type == LayoutType.feature_grid and final.exact_element_count == 4
    assert final.story_stage == "industry impact"

def test_structured_stress_brief_keeps_native_table_and_five_step_roadmap(tmp_path):
    prompt='''Create an intensive 6-slide deck. Slide 1: Title Slide. The title must be a long phrase: "Project Hyperion: Scaling Global Microservices Infrastructure". Slide 2: The Core Issue. Summarize this into 3 distinct detailed bullet points. Slide 3: System Component Comparison. Generate a 4x4 markdown table. Columns: [Service Name, Current Latency (ms), Target Latency (ms), Risk Level]. Row 1: [AuthGate API Gateway, 450ms, <15ms, Critical Risk / High Priority]. Row 2: [DataStream Ledger Sync, 1,200ms, <50ms, High Risk / Complex Migration]. Row 3: [NotifyEngine PubSub, 85ms, <10ms, Low Risk / Fast Win]. Slide 4: Migration Timeline. A horizontal 5-step engineering roadmap sequence. Step 1: Discovery & Audit, Step 2: Protocol Definition & RFC, Step 3: Canary Deployments in Sandbox, Step 4: Multi-Region Traffic Cutover, Step 5: Legacy Decommissioning & Cleanup. Slide 5: Critical Metrics. Display 3 distinct large percentage metrics side-by-side: 99.999% (Label: Targeted Uptime SLA), -85% (Label: Reduction in P99 API Latency), and $4.2M (Label: Projected Annual Savings). Slide 6: Emergency Contacts. Include an On-Call Matrix with varying contact lengths.'''
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=6)
    assert len(brief.slides) == 6
    assert brief.by_number(1).title.startswith("Project Hyperion")
    table_slide=brief.by_number(3)
    assert table_slide.table_data["headers"] == ["Service Name", "Current Latency (ms)", "Target Latency (ms)", "Risk Level"]
    assert table_slide.table_data["rows"][1][1] == "1,200ms"
    roadmap=brief.by_number(4)
    assert roadmap.exact_element_count == 5 and len(roadmap.requirements) == 5
    metrics=brief.by_number(5)
    assert [item.value for item in metrics.elements] == ["99.999%", "-85%", "$4.2M"]
    spec=PresentationSpec(title="Project Hyperion", topic="Project Hyperion", slides=[item.seed() for item in brief.slides])
    PresentationQAAgent().validate_and_recompose(spec)
    assert len(spec.slides[3].elements) == 5
    output=build_presentation(spec,tmp_path/"stress.pptx")
    rendered=Presentation(output)
    assert len(rendered.slides[2].shapes) > 3
    assert any(getattr(shape, "has_table", False) for shape in rendered.slides[2].shapes)

def test_metric_echoes_and_dangling_sentence_are_removed():
    spec=PresentationSpec(title="Metrics", topic="Metrics", slides=[
        SlideSpec(slide_number=1, title="Roadmap", purpose="Show the roadmap.", layout_type=LayoutType.step_workflow,
                  elements=[{"heading":"Cleanup", "body":"Final shutdown of old infrastructure and removal of redundant code paths to optimize."}]),
        SlideSpec(slide_number=2, title="Metrics", purpose="Show the metrics.", layout_type=LayoutType.key_metrics,
                  elements=[{"heading":"Targeted Uptime SLA", "value":"99.999%", "body":"Targeted Uptime SLA"}, {"heading":"Latency reduction", "value":"-85%", "body":"-85%"}]),
    ])
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].elements[0].body.endswith("code paths.")
    assert all(not item.body for item in spec.slides[1].elements)
    assert summarize_point("Higher Conversion Rates vs. Standard Web", 8) == "Higher Conversion Rates versus Standard Web"

def test_metric_renderer_never_draws_duplicate_value_body(tmp_path):
    spec=PresentationSpec(title="Metrics", topic="Metrics", slides=[SlideSpec(
        slide_number=1, title="Growth", purpose="Growth evidence.", layout_type=LayoutType.key_metrics,
        elements=[{"heading":"Higher Conversion Rates vs. Standard Web", "value":"3.8x", "body":"3.8x"}],
    )])
    output=build_presentation(spec,tmp_path/"metrics-no-echo.pptx")
    texts=[shape.text for shape in Presentation(output).slides[0].shapes if getattr(shape,"has_text_frame",False)]
    assert texts.count("3.8x") == 1

def test_explicit_metric_contract_cannot_be_overwritten_with_an_echo_value():
    brief=BriefInterpreterAgent().interpret(
        'Slide 1: Critical Metrics\nLayout: Key Metrics\nDisplay metrics: 99.999% (Label: Targeted Uptime SLA).', 3
    )
    directive=brief.by_number(1)
    preserved=SlideContentAgent._preserve_contract_elements(directive, {
        "elements":[{"heading":"Targeted Uptime SLA", "value":"99.999%", "body":"99.999%"}]
    })
    assert preserved[0]["value"] == "99.999%"
    assert preserved[0]["heading"] == "Targeted Uptime SLA"
    assert not preserved[0]["body"]

def test_slide_content_prompts_do_not_repeat_other_slide_contracts(monkeypatch):
    prompt='''Create a 3-slide deck titled "Scoped prompts".
Slide 1: Cover
Title: Opening signal
Slide 2: Private marker
Layout: Feature Grid
Topic Areas: NEVER-SEND-THIS-TO-SLIDE-ONE
Slide 3: Closing decision
Layout: Summary'''
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=3)
    sent=[]

    class FakeGateway:
        def generate_json(self, payload, **_kwargs):
            sent.append(payload)
            return {
                "title":"Generated slide", "subtitle":None, "layout_type":"feature_grid",
                "purpose":"A complete, audience-facing insight.", "elements":[], "visual_spec":{},
            }

    monkeypatch.setattr(
        "app.agents.slide_content_agent.LLMGateway.from_settings",
        lambda *_args, **_kwargs: FakeGateway(),
    )
    SlideContentAgent().generate(
        CreatePresentationRequest(topic=prompt, slide_count=3), "Cyber Dark", StorylineAgent(),
        provider="gemini", brief=brief,
    )
    assert len(sent) == 3
    assert "NEVER-SEND-THIS-TO-SLIDE-ONE" not in sent[0]
    assert "NEVER-SEND-THIS-TO-SLIDE-ONE" in sent[1]
    assert "Slide 2: Private marker" not in sent[0]

def test_portable_gradient_background_and_long_metric_copy(tmp_path):
    spec=PresentationSpec(
        title="Cost comparison", topic="Compare costs",
        design_system={"name":"Corporate Blue", "background_color":"#FFFFFF", "surface_color":"#F7F9FC", "primary_color":"#2167C5", "secondary_color":"#6B96D0", "header_color":"#1A2D45", "text_primary":"#1A2D45", "text_secondary":"#536476", "muted_text":"#6D8092"},
        slides=[SlideSpec(
            slide_number=1, title="Cost and Efficiency Comparison", purpose="Compare practical costs", layout_type=LayoutType.key_metrics,
            visual_spec={"background_treatment":"halo"},
            elements=[
                {"heading":"NSDL: The Institutional Standard", "body":"Higher focus on institutional stability with robust security protocols for large-scale holdings."},
                {"heading":"CDSL: The Retail Favorite", "body":"Lower transaction costs and seamless digital integration for active individual traders."},
                {"heading":"Key Cost Driver", "body":"Broker and transaction charges determine the practical cost for most retail investors."},
            ],
        )],
    )
    output=build_presentation(spec,tmp_path/"portable-background.pptx")
    presentation=Presentation(output)
    text_boxes=[shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    heading=next(shape for shape in text_boxes if shape.text=="NSDL: The Institutional Standard")
    body=next(shape for shape in text_boxes if shape.text.startswith("Higher focus"))
    assert heading.top+heading.height < body.top
    with zipfile.ZipFile(output) as archive:
        assert b"a:gradFill" in archive.read("ppt/slides/slide1.xml")

def test_asymmetric_insight_uses_one_page_number_and_fits_card_body(tmp_path):
    spec=PresentationSpec(
        title="Depository choices", topic="Compare depositories",
        slides=[SlideSpec(
            slide_number=1, title="The Infrastructure of Your Wealth", purpose="Establish that the choice between NSDL and CDSL is the foundational decision.", layout_type=LayoutType.content_with_visual,
            elements=[
                {"heading":"The Regulatory Landscape", "body":"Two distinct central depositories manage your securities, ensuring safe electronic transfers."},
                {"heading":"Why the Choice Matters", "body":"Your choice affects broker compatibility and the platform experience."},
            ],
        )],
    )
    output=build_presentation(spec,tmp_path/"asymmetric.pptx")
    presentation=Presentation(output)
    text_boxes=[shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    page_numbers=[shape for shape in text_boxes if shape.text=="1"]
    body=next(shape for shape in text_boxes if shape.text.startswith("Two distinct"))
    assert len(page_numbers)==1
    assert body.top+body.height <= 3.67 * 914400

def test_storyline_uses_distinct_security_and_investor_arcs():
    def plan(topic):
        spec=PresentationSpec(title=topic, topic=topic, slides=[
            SlideSpec(slide_number=index, title=f"Slide {index}", purpose="A concrete point.", layout_type=LayoutType.feature_grid)
            for index in range(1, 7)
        ])
        return StorylineAgent().apply(spec)
    security=plan("Autonomous threat mitigation and SOC orchestration")
    investor=plan("NSDL vs CDSL: which depository should I choose?")
    assert security.arc == "security operations arc"
    assert investor.arc == "investor decision arc"
    assert [beat.preferred_recipe for beat in security.beats[1:-1]] != [beat.preferred_recipe for beat in investor.beats[1:-1]]

def test_auto_theme_is_topic_aware():
    assert auto_theme_for_topic("SOC incident response automation").name == "Security Signal"
    assert auto_theme_for_topic("NSDL versus CDSL investor choice").name == "Investor Slate"

def test_explicit_light_brand_prompt_overrides_auto_topic_theme():
    theme=resolve_theme(
        "Auto",
        "Use a clean Light Theme with an off-white background. Do not use dark backgrounds. Brand colors are Deep Navy Blue, Warm Coral/Terracotta, and Charcoal Gray.",
    )
    assert theme.name == "Luxury Light Brand"
    assert theme.background_color == "#FBF7F1"
    assert theme.header_color == "#102A43"
    assert theme.accent_color == "#C96B5A"
    assert theme.font_heading == "Georgia"

def test_light_brand_brief_keeps_deck_title_quoted_metrics_and_split_layout():
    prompt='''Create a 5-slide presentation titled "The Future of Luxury Retail: Spatial Commerce & AR". Use a Light Theme; do not use dark backgrounds. Slide 1: Title Slide. Slide 2: Market Shift. Create a split-screen 2-column layout. Slide 3: Growth Projections. Metrics: "+142%" (Label: Increase in Customer Dwell Time), "$12B" (Label: AR Retail Market Value by 2028), and "3.8x" (Label: Higher Conversion Rates vs. Standard Web). Slide 4: Pillars. Slide 5: Partnerships. Display 4 separate items.'''
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=5)
    assert brief.by_number(1).title == "The Future of Luxury Retail: Spatial Commerce & AR"
    assert brief.by_number(2).layout_type == LayoutType.two_column
    metrics=brief.by_number(3)
    assert metrics.layout_type == LayoutType.key_metrics
    assert [item.value for item in metrics.elements] == ["+142%", "$12B", "3.8x"]

def test_dense_devops_brief_preserves_tiers_nodes_and_quadrant_count():
    prompt='''Create a 5-slide deck titled "DevOps". Slide 1: Title Slide. Main Title: "Next-Generation DevOps: Enterprise CI/CD". Slide 2: Core Architecture. Create 3 main architectural tiers with nested sub-bullets:Automated Build & Compilation LayerTriggered by webhook events from Version Control Systems.Executes builds in Kubernetes pods.Distributed Continuous Testing MatrixParallel test execution dynamically.Progressive Deployment EngineCanary releases with rollback triggers.Slide 3: Variable Data Rows. Node A: Build. (Extremely short)Node B: Code Quality Scan. (Medium)Node C: Multi-Region Distributed Compliance Verification Module. (Extremely long phrase)Node D: Deploy. (Extremely short)Slide 4: Pillars. A vertical list with 3 pillars. Slide 5: Recovery. Display a 2x2 grid representing failover clusters.'''
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=5)
    assert brief.by_number(1).title == "Next-Generation DevOps: Enterprise CI/CD"
    tiers=brief.by_number(2)
    assert tiers.layout_type == LayoutType.architecture_layers and tiers.nested_bullets and len(tiers.elements) == 3
    assert "Triggered by webhook events" in tiers.elements[0].body
    assert "Executes builds in Kubernetes pods" in tiers.elements[0].body
    nodes=brief.by_number(3)
    assert nodes.layout_type == LayoutType.comparison and nodes.variable_rows and len(nodes.elements) == 4
    assert nodes.elements[1].heading == "NODE B: Code Quality Scan"
    assert not nodes.elements[2].body
    assert brief.by_number(4).exact_element_count == 3
    assert brief.by_number(5).exact_element_count == 4

def test_qa_preserves_explicit_long_title_and_nested_source_copy(tmp_path):
    brief=BriefInterpreterAgent().interpret(
        '''Create a 3-slide deck titled "Pipeline". Slide 1: Title Slide. Main Title: "Next-Generation DevOps: Architecting Scalable, Zero-Downtime CI/CD Pipelines for Enterprise Infrastructure v3.0". Slide 2: Architecture. Create 3 main architectural tiers with nested sub-bullets:Automated Build LayerTriggered by webhooks.Executes isolated builds.Distributed Testing MatrixRuns tests in parallel.Progressive Deployment EngineCanary releases with rollback. Slide 3: Close.''',
        3,
    )
    spec=PresentationSpec(title="Pipeline", topic="Pipeline", slides=[item.seed() for item in brief.slides])
    PresentationQAAgent().validate_and_recompose(spec)
    assert spec.slides[0].title.endswith("Infrastructure v3.0")
    assert "Triggered by webhooks." in spec.slides[1].elements[0].body
    assert "Executes isolated builds." in spec.slides[1].elements[0].body
    output=build_presentation(spec, tmp_path/"nested-source-copy.pptx")
    rendered_text=[shape.text for shape in Presentation(output).slides[1].shapes if getattr(shape, "has_text_frame", False)]
    assert any("Triggered by webhooks." in text and "Executes isolated builds." in text for text in rendered_text)

def test_supply_chain_brief_keeps_generic_nested_divisions_phases_and_light_theme():
    prompt='''Create a 5-slide presentation titled "Global Supply Chain Optimization: Resilience Framework 2027". Slide 1: Title Slide. Use a modern, light-themed abstract globe background with crisp lines. Slide 2: Strategic Pillars. Create 3 core operational divisions:Multi-Tier Supplier DiversificationMapping tier-1 through tier-3 dependency bottlenecks.Establishing secondary sourcing redundancy pipelines.Predictive Logistical TelemetryIntegrating IoT sensors to track route disruptions.Dynamic Inventory BufferingDeploying near-shore warehousing strategies. Slide 3: Regional Risk Analysis. Columns: [Region, Risk, Bottleneck, Mitigation] Row 1: [APAC, 8.7 / 10, Congestion, Alternative feeder networks] Slide 4: Metrics. Slide 5: Execution Roadmap. Display a 4-step horizontal timeline: Phase 1: Global Risk Audit, Phase 2: System API Integration, Phase 3: Pilot Node Deployment, Phase 4: Full Network Cutover.'''
    brief=BriefInterpreterAgent().interpret(prompt, 5)
    divisions=brief.by_number(2)
    assert divisions.layout_type == LayoutType.architecture_layers and divisions.nested_bullets
    assert [item.heading for item in divisions.elements] == ["Multi-Tier Supplier Diversification", "Predictive Logistical Telemetry", "Dynamic Inventory Buffering"]
    assert "secondary sourcing redundancy pipelines" in divisions.elements[0].body
    roadmap=brief.by_number(5)
    assert roadmap.layout_type == LayoutType.step_workflow
    assert roadmap.requirements == ["Global Risk Audit", "System API Integration", "Pilot Node Deployment", "Full Network Cutover"]
    assert resolve_theme("Auto", prompt).name == "Modern Light"

def test_slide_edit_honors_horizontal_nested_division_instruction(tmp_path):
    spec=PresentationSpec(title="Supply chain", topic="Supply chain", slides=[
        SlideSpec(slide_number=1, title="Strategic Pillars", purpose="Current content.", layout_type=LayoutType.architecture_layers),
    ])
    instruction="""Create a dense horizantal layout featuring 3 core operational divisions:Multi-Tier Supplier DiversificationMapping tier-1 through tier-3 bottlenecks.Establishing secondary sourcing redundancy.Predictive Logistical TelemetryIntegrating IoT sensors to track disruptions.Dynamic Inventory BufferingDeploying near-shore warehousing strategies."""
    edited=PresentationOrchestrator().edit_slide(spec, 1, instruction)
    slide=edited.slides[0]
    assert slide.layout_type == LayoutType.feature_grid
    assert slide.visual_spec["horizontal_nested"] is True
    assert [item.heading for item in slide.elements] == ["Multi-Tier Supplier Diversification", "Predictive Logistical Telemetry", "Dynamic Inventory Buffering"]
    output=build_presentation(edited, tmp_path/"horizontal-nested-edit.pptx")
    text=[shape.text for shape in Presentation(output).slides[0].shapes if getattr(shape, "has_text_frame", False)]
    assert any("tier-1 through tier-3" in value for value in text)

def test_slide_edit_relayouts_existing_content_and_accepts_image_and_font_changes():
    spec=PresentationSpec(title="Deck", topic="Deck", slides=[SlideSpec(
        slide_number=1, title="Existing slide", purpose="Keep the current content.", layout_type=LayoutType.feature_grid,
        elements=[{"heading":"One","body":"First"},{"heading":"Two","body":"Second"},{"heading":"Three","body":"Third"}],
    )])
    edited=PresentationOrchestrator().edit_slide(
        spec, 1, "Change structure form 3 to column. Keep existing data. Add an image of a connected logistics network and set font size to 18pt.",
    )
    slide=edited.slides[0]
    assert [item.heading for item in slide.elements] == ["One", "Two", "Three"]
    assert slide.visual_spec["grid_columns"] == 3
    assert slide.visual_spec["image_required"] is True
    assert slide.visual_spec["edit_image_requested"] is True
    assert slide.visual_spec["font_scale"] > 1

def test_minimalist_and_corporate_themes_are_visually_distinct():
    from app.agents.orchestrator import THEMES
    minimalist=THEMES["Minimalist White"]
    corporate=THEMES["Corporate Blue"]
    assert minimalist.background_color != corporate.background_color
    assert minimalist.primary_color != corporate.primary_color
    assert minimalist.header_color != corporate.header_color

def test_unsplash_status_explains_a_missing_server_key():
    service=ImageService()
    service.settings=Settings(image_provider="unsplash", unsplash_access_key="")
    status=service.status()
    assert status["provider"] == "unsplash"
    assert status["ready"] is False
    assert "UNSPLASH_ACCESS_KEY" in status["message"]

def test_generation_progress_names_active_agents():
    events=[]
    PresentationOrchestrator().generate(
        CreatePresentationRequest(topic="A platform strategy", slide_count=3),
        progress=lambda stage, percent: events.append((stage, percent)),
    )
    labels=[stage for stage, _ in events]
    assert any("Brief Classification Agent" in label for label in labels)
    assert any("Design Director Agent" in label for label in labels)
    assert any("Presentation QA Agent" in label for label in labels)

def test_brief_interpreter_preserves_explicit_slide_contracts_in_fallback():
    prompt='''Agentic AI Alert Response Platform
Slide 1: Title Slide
Title: Autonomous Alert Triage Platform
Subtitle: Transforming SOC Operations with Agentic Incident Response
Slide 2: Core Capabilities
Layout: Feature Grid (4 items)
Topic Areas:
Autonomous Context Enrichment (correlating telemetry and indicators)
Dynamic Investigation Plans (building hypothesis-driven triage paths)
Human-in-the-Loop Safeguards (requiring analyst approval for high-risk actions)
Self-Healing Playbooks (updating detections from retrospectives)
Slide 3: Incident Response Pipeline
Layout: Step Workflow (4 sequential steps)
Steps to cover:
Alert Ingestion (streaming SIEM and EDR data)
Agent Triage (evaluating severity)
Action Execution (isolating endpoints)
Post-Mortem Audit (updating tickets)'''
    brief=BriefInterpreterAgent().interpret(prompt, 3)
    assert brief.is_structured and brief.by_number(2).layout_type == LayoutType.feature_grid
    assert len(brief.by_number(2).requirements) == 4
    # The UI can still be on its six-slide default; explicit Slide 1–3
    # instructions are the authoritative deck contract.
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic=prompt, slide_count=6))
    assert len(spec.slides) == 3
    assert spec.slides[0].title == "Autonomous Alert Triage Platform"
    assert spec.slides[1].layout_type == LayoutType.feature_grid
    assert len(spec.slides[1].elements) == 4
    assert spec.slides[2].layout_type == LayoutType.step_workflow

def test_design_director_keeps_explicit_feature_grid_recipe():
    spec=PresentationSpec(title="Impact", topic="SOC impact", slides=[
        SlideSpec(slide_number=1, title="Cover", purpose="Start", layout_type=LayoutType.title_slide),
        SlideSpec(slide_number=2, title="Three metrics", purpose="Evidence", layout_type=LayoutType.feature_grid, metadata={"brief_layout_locked":True}),
        SlideSpec(slide_number=3, title="Locked metrics", purpose="End", layout_type=LayoutType.feature_grid, metadata={"brief_layout_locked":True}),
    ])
    DesignDirectorAgent().apply(spec)
    assert spec.slides[1].visual_spec["composition"] == "editorial insight grid"
    assert spec.slides[2].visual_spec["composition"] == "editorial insight grid"

def test_brief_classifier_distinguishes_structured_and_open_ended_prompts():
    agent=BriefInterpreterAgent()
    structured=agent.classify("Slide 1: Cover\nTitle: A platform\nSlide 2: Workflow\nLayout: Step Workflow")
    generic=agent.classify("Explain how an agentic AI security platform improves SOC operations.")
    assert structured.mode == "structured"
    assert structured.detected_slide_numbers == [1, 2]
    assert generic.mode == "open_ended"

def test_brief_classifier_recognizes_markdown_slide_contracts():
    prompt='''**Slide 1: Title Slide**
- Title: Autonomous Alert Triage Platform
- Subtitle: Transforming SOC Operations

**Slide 2: Core Capabilities**
- Layout: Feature Grid (4 items)
- Topic Areas:
1. Context Enrichment (correlate telemetry)
2. Investigation Plans (triage paths)
3. Safeguards (analyst approval)
4. Playbooks (retrospectives)'''
    agent=BriefInterpreterAgent()
    brief=agent.interpret(prompt, 6)
    assert agent.classify(prompt).mode == "structured"
    assert brief.by_number(1).title == "Autonomous Alert Triage Platform"
    assert brief.by_number(2).layout_type == LayoutType.feature_grid
    assert brief.by_number(2).requirements == [
        "Context Enrichment (correlate telemetry)", "Investigation Plans (triage paths)",
        "Safeguards (analyst approval)", "Playbooks (retrospectives)",
    ]

def test_brief_interpreter_preserves_compact_prose_slide_agenda():
    """A mobile/copy-paste prompt may lose every line break after `Slides:`."""
    prompt=(
        'Create a 10-slide presentation titled "Enterprise Agentic AI Platform Architecture". '
        'Slides: Title slide with a concise value proposition. '
        'Business problem — fragmented enterprise knowledge, manual workflows, slow decisions, and hallucination risks. '
        'Proposed solution — show a high-level architecture diagram from users → AI application → agent orchestration → RAG → LLM → enterprise systems. '
        'Detailed RAG pipeline showing ingestion, document processing, chunking, embeddings, vector storage, retrieval, reranking, prompt construction, and generation. '
        'Agentic workflow showing planner, specialized agents, tool calling, memory, validation, and human approval. '
        'Azure-based deployment architecture using API Management, App Services/AKS, Azure Functions, Service Bus, Blob Storage, Azure AI Search, Key Vault, Application Insights, and Azure OpenAI. '
        'Security architecture covering Managed Identity, RBAC, network isolation, secrets management, PII protection, prompt injection protection, and audit logging. '
        'Scalability and reliability architecture covering horizontal scaling, queues, caching, circuit breakers, retries, dead-letter queues, and observability. '
        'Cost optimization table comparing major infrastructure components and optimization strategies. '
        'Implementation roadmap divided into MVP, production hardening, enterprise rollout, and autonomous-agent phase. '
        'Use native editable PowerPoint shapes and connectors.'
    )
    agent=BriefInterpreterAgent()
    brief=agent.interpret(prompt, requested_count=10)
    assert agent.classify(prompt).mode == "structured"
    assert brief.deck_title == "Enterprise Agentic AI Platform Architecture"
    assert len(brief.slides) == 10
    assert brief.by_number(3).layout_type == LayoutType.process_flow
    assert [item.heading for item in brief.by_number(3).elements] == [
        "Users", "AI application", "Agent orchestration", "RAG", "LLM", "Enterprise systems",
    ]
    assert brief.by_number(6).native_diagram is True

def test_brief_contract_overrides_the_default_slide_slider_up_to_ten():
    prompt="\n".join(
        f"Slide {number}: Slide {number} title\nLayout: {'Dashboard' if number == 8 else 'Feature Grid'}"
        for number in range(1, 11)
    )
    brief=BriefInterpreterAgent().interpret(prompt, requested_count=6)
    assert len(brief.slides) == 10
    assert brief.by_number(8).layout_type == LayoutType.dashboard
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic=prompt, slide_count=6))
    assert len(spec.slides) == 10

def test_brief_interpreter_protects_embedded_json_slide_schema():
    prompt='''Use this JSON schema:
    {"title":"Agentic AI Alert Response Platform","slides":[
      {"slide_number":1,"title":"Autonomous Alert Triage Platform","subtitle":"Transforming SOC Operations","layout_type":"title_slide","elements":[]},
      {"slide_number":2,"title":"Platform Core Capabilities","layout_type":"feature_grid","elements":[
        {"id":"e1","heading":"Context Enrichment","subtext":"Correlate SIEM telemetry."},
        {"id":"e2","heading":"Investigation Plans","subtext":"Generate triage paths."},
        {"id":"e3","heading":"Approval Safeguards","subtext":"Gate high-risk actions."},
        {"id":"e4","heading":"Self-Healing Playbooks","subtext":"Learn from incidents."}]},
      {"slide_number":3,"title":"Response Pipeline","layout_type":"step_workflow","elements":[]},
      {"slide_number":4,"title":"Architecture Layers","layout_type":"architecture_layers","elements":[]},
      {"slide_number":5,"title":"Strategic SOC Impact","layout_type":"feature_grid","elements":[]}
    ]}'''
    agent=BriefInterpreterAgent()
    brief=agent.interpret(prompt, 6)
    assert agent.classify(prompt).mode == "structured"
    assert brief.deck_title == "Agentic AI Alert Response Platform"
    assert len(brief.slides) == 5
    assert brief.by_number(2).elements[0].body == "Correlate SIEM telemetry."
    spec=PresentationOrchestrator().generate(CreatePresentationRequest(topic=prompt, slide_count=6))
    assert len(spec.slides) == 5
    assert spec.title == "Agentic AI Alert Response Platform"
    assert len(spec.slides[1].elements) == 4

def test_slide_content_contract_restores_items_missing_from_model_response():
    brief=BriefInterpreterAgent().interpret('''{"slides":[{"slide_number":1,"title":"Capabilities","layout_type":"feature_grid","elements":[{"heading":"Context","subtext":"Correlate telemetry."},{"heading":"Containment","subtext":"Isolate endpoints."}]}]}''', 1)
    directive=brief.by_number(1)
    restored=SlideContentAgent._preserve_contract_elements(directive, {"elements":[{"heading":"Context","body":"Model-expanded context."}]})
    assert [item["heading"] for item in restored] == ["Context", "Containment"]
    assert restored[0]["body"] == "Model-expanded context."

def test_slide_content_repairs_a_component_fragment_before_slide_validation():
    source=SlideSpec(slide_number=2, title="Fallback title", purpose="Fallback purpose.", layout_type=LayoutType.feature_grid)
    payload=SlideContentAgent._complete_slide_payload(
        source,
        None,
        {"type":"comparison_group", "title":None, "elements":{"left":"Manual", "right":"AI"}},
    )
    repaired=SlideSpec.model_validate(payload)
    assert repaired.title == "Fallback title"
    assert repaired.purpose == "Fallback purpose."
    assert repaired.layout_type == LayoutType.feature_grid
    assert repaired.metadata["provider_response_repaired"] == ["title", "purpose", "elements", "visual_spec"]

def test_web_preview_matches_pptx_two_by_two_feature_grid():
    spec=PresentationSpec(title="Capabilities", topic="SOC", slides=[SlideSpec(
        slide_number=1, title="Core Capabilities", purpose="Show the platform pillars.", layout_type=LayoutType.feature_grid,
        elements=[{"heading":f"Capability {index}", "body":"Technical detail."} for index in range(1, 5)],
    )])
    rendered=render_slide_html(spec, 1)
    assert "grid-template-columns:repeat(2,1fr)" in rendered
    assert "max-width:1120px" in rendered

def test_chevron_flow_keeps_step_explanations_inside_connected_stages(tmp_path):
    spec=PresentationSpec(title="Workflow", topic="Workflow", slides=[SlideSpec(
        slide_number=1, title="Generation workflow", purpose="Show the sequence.", layout_type=LayoutType.process_flow,
        visual_spec={"visual_variant":"chevron_flow"},
        elements=[{"heading":f"Step {index}", "body":"A concise explanation that belongs to this stage."} for index in range(1, 5)],
    )])
    output=build_presentation(spec,tmp_path/"workflow.pptx")
    presentation=Presentation(output)
    assert all(shape.text != "DETAIL" for shape in presentation.slides[0].shapes if shape.has_text_frame)
    rendered=render_slide_html(spec, 1)
    assert "flow-detail" not in rendered

def test_standard_diagrams_keep_copy_inside_shapes_in_both_renderers(tmp_path):
    elements=[{"heading":f"Layer {index}", "body":"A concise explanation contained by the diagram shape."} for index in range(1, 4)]
    for layout in (LayoutType.architecture_layers, LayoutType.process_flow):
        spec=PresentationSpec(title="Diagrams", topic="Diagrams", slides=[SlideSpec(
            slide_number=1, title="System diagram", purpose="Show the system.", layout_type=layout, elements=elements,
        )])
        output=build_presentation(spec,tmp_path/f"{layout.value}.pptx")
        presentation=Presentation(output)
        body_shapes=[shape for shape in presentation.slides[0].shapes if shape.has_text_frame and shape.text.startswith("A concise explanation")]
        assert len(body_shapes) == 3
        rendered=render_slide_html(spec, 1)
        expected="architecture-layers" if layout == LayoutType.architecture_layers else "standard-flow"
        assert expected in rendered

def test_component_animations_are_native_and_capped(tmp_path):
    spec=PresentationSpec(title="Response", topic="SOC", slides=[SlideSpec(
        slide_number=1, title="Response workflow", purpose="Show stages.", layout_type=LayoutType.process_flow,
        visual_spec={"story_stage":"response lifecycle"},
        elements=[{"heading":f"Stage {index}", "body":"A complete explanation for the stage."} for index in range(1, 5)],
    )])
    output=build_presentation(spec,tmp_path/"animated.pptx")
    with zipfile.ZipFile(output) as archive:
        xml=archive.read("ppt/slides/slide1.xml")
    assert b"<p:timing>" in xml
    assert xml.count(b"<p:animEffect") <= 7  # title + three content groups, each with heading/body
    assert b'filter="wipe(left)"' in xml

def test_source_backed_chart_is_native_in_pptx_and_web_preview(tmp_path):
    spec=PresentationSpec(title="Response time", topic="SOC", slides=[SlideSpec(
        slide_number=1, title="MTTR trend", purpose="Source-backed response time trend.", layout_type=LayoutType.key_metrics,
        visual_spec={"chart_data":{"type":"column","categories":["Q1","Q2","Q3"],"series":[{"name":"MTTR minutes","values":[30,18,9]}]}},
    )])
    output=build_presentation(spec,tmp_path/"chart.pptx")
    presentation=Presentation(output)
    assert any(shape.has_chart for shape in presentation.slides[0].shapes)
    assert "native-chart-preview" in render_slide_html(spec, 1)

def test_chevron_flow_places_step_details_inside_connected_stages(tmp_path):
    spec=PresentationSpec(title="Pipeline", topic="SOC", slides=[SlideSpec(
        slide_number=1, title="Incident Response Pipeline", purpose="Show the response sequence.", layout_type=LayoutType.step_workflow,
        visual_spec={"visual_variant":"chevron_flow"},
        elements=[{"heading":f"Step {index}", "body":"Detailed operational explanation for this response stage."} for index in range(1,5)],
    )])
    output=build_presentation(spec,tmp_path/"flow.pptx")
    presentation=Presentation(output)
    bodies=[shape for shape in presentation.slides[0].shapes if shape.has_text_frame and shape.text.startswith("Detailed operational")]
    assert len(bodies) == 4
    assert all(shape.top / 914400 >= 4.0 and (shape.top+shape.height) / 914400 <= 5.5 for shape in bodies)
    assert "flow-detail" not in render_slide_html(spec, 1)

def test_request_accepts_a_detailed_structured_brief():
    request=CreatePresentationRequest(topic="{" + '"slides":[],' * 700 + "}")
    assert len(request.topic) > 2_000
