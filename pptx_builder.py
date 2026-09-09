from __future__ import annotations

from io import BytesIO
from typing import Any, Dict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from llm_service import HighlightState, LayoutType, PresentationSpec, SlideSpec

THEMES: Dict[str, Dict[str, Any]] = {
    "Cyber Dark": {
        "background": "0B1320",
        "surface": "142136",
        "surface_2": "1D2D47",
        "accent": "00E5FF",  # High-vis Cyan
        "active": "FF9100",  # High-vis Amber
        "header": "FFFFFF",  # Pure White
        "text": "E2E8F0",    # Light Slate Text
        "muted": "94A3B8",   # Soft Gray Subtitles
        "warning": "FF5252",
        "line": "2A3B56",
    },
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 14,
    color: str = "FFFFFF",
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font_name: str = "Arial",
    margin: float = 0.04,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _add_rect(slide, x: float, y: float, w: float, h: float, *, fill: str, line: str | None = None, radius: bool = True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape


def _add_header(slide, spec: SlideSpec, theme: Dict[str, Any]):
    _add_text(slide, f"{spec.slide_number:02d}", 0.55, 0.35, 0.55, 0.35, size=12, color=theme["accent"], bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, spec.title, 1.25, 0.25, 11.3, 0.55, size=22, color=theme["header"], bold=True)
    if spec.subtitle:
        _add_text(slide, spec.subtitle, 1.25, 0.82, 11.3, 0.42, size=12, color=theme["muted"])

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.32), Inches(12.25), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = _rgb(theme["line"])
    line.line.fill.background()


def _element_colors(element, theme):
    if element.highlight_state == HighlightState.active:
        return theme["active"], theme["surface_2"]
    if element.highlight_state == HighlightState.warning:
        return theme["warning"], theme["surface_2"]
    return theme["accent"], theme["surface"]


def _render_title(slide, spec: SlideSpec, theme: Dict[str, Any]):
    _add_text(slide, spec.title, 0.8, 2.0, 11.8, 1.2, size=34, color=theme["header"], bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, spec.subtitle or "Executive Strategic Deck", 1.3, 3.3, 10.8, 0.7, size=16, color=theme["accent"], align=PP_ALIGN.CENTER)
    _add_rect(slide, 4.1, 4.3, 5.15, 0.05, fill=theme["accent"], line=None, radius=False)
    _add_text(slide, "AI PRESENTATION ENGINE", 4.1, 4.5, 5.15, 0.35, size=10, color=theme["muted"], bold=True, align=PP_ALIGN.CENTER)


def _render_grid(slide, spec: SlideSpec, theme: Dict[str, Any], columns: int = 2):
    elements = spec.elements[:8]
    if not elements:
        return
    rows = (len(elements) + columns - 1) // columns
    gap_x, gap_y = 0.35, 0.3
    start_x, start_y = 0.65, 1.7
    card_w = (12.0 - gap_x * (columns - 1)) / columns
    card_h = (5.15 - gap_y * (rows - 1)) / rows

    for idx, element in enumerate(elements):
        row, col = divmod(idx, columns)
        x = start_x + col * (card_w + gap_x)
        y = start_y + row * (card_h + gap_y)
        accent, fill = _element_colors(element, theme)

        _add_rect(slide, x, y, card_w, card_h, fill=fill, line=theme["line"])
        _add_rect(slide, x, y, 0.08, card_h, fill=accent, line=None, radius=False)
        _add_text(slide, element.heading, x + 0.25, y + 0.2, card_w - 0.4, 0.4, size=16, color=theme["header"], bold=True)
        _add_text(slide, element.subtext, x + 0.25, y + 0.7, card_w - 0.4, card_h - 0.8, size=12, color=theme["text"])


def _render_workflow(slide, spec: SlideSpec, theme: Dict[str, Any]):
    elements = spec.elements[:7]
    if not elements:
        return
    start_x, y = 0.65, 2.8
    gap, node_w, node_h = 0.15, 1.62, 2.2

    for idx, element in enumerate(elements):
        x = start_x + idx * (node_w + gap)
        accent, fill = _element_colors(element, theme)

        _add_rect(slide, x, y, node_w, node_h, fill=fill, line=accent)
        _add_text(slide, str(idx + 1).zfill(2), x + 0.15, y + 0.15, 0.5, 0.3, size=12, color=accent, bold=True)
        _add_text(slide, element.heading, x + 0.15, y + 0.55, node_w - 0.3, 0.5, size=13, color=theme["header"], bold=True)
        _add_text(slide, element.subtext, x + 0.15, y + 1.1, node_w - 0.3, 1.0, size=10, color=theme["text"])

        if idx < len(elements) - 1:
            _add_rect(slide, x + node_w, y + node_h / 2 - 0.02, gap, 0.04, fill=theme["accent"], line=None, radius=False)


def _render_layers(slide, spec: SlideSpec, theme: Dict[str, Any]):
    elements = spec.elements[:6]
    if not elements:
        return
    start_y, layer_h, gap = 1.75, 0.8, 0.15
    for idx, element in enumerate(elements):
        y = start_y + idx * (layer_h + gap)
        accent, fill = _element_colors(element, theme)

        _add_rect(slide, 0.8, y, 11.7, layer_h, fill=fill, line=theme["line"])
        _add_rect(slide, 0.8, y, 0.12, layer_h, fill=accent, line=None, radius=False)
        _add_text(slide, f"L{idx + 1}", 1.1, y + 0.2, 0.5, 0.4, size=12, color=accent, bold=True)
        _add_text(slide, element.heading, 1.7, y + 0.18, 3.5, 0.4, size=15, color=theme["header"], bold=True)
        _add_text(slide, element.subtext, 5.3, y + 0.18, 7.0, 0.45, size=11, color=theme["text"])


def build_presentation(spec: PresentationSpec, theme_name: str = "Cyber Dark") -> bytes:
    theme = THEMES.get(theme_name, THEMES["Cyber Dark"])
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for spec_slide in spec.slides:
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(theme["background"])

        if spec_slide.layout_type == LayoutType.title_slide:
            _render_title(slide, spec_slide, theme)
        else:
            _add_header(slide, spec_slide, theme)
            if spec_slide.layout_type == LayoutType.feature_grid:
                _render_grid(slide, spec_slide, theme)
            elif spec_slide.layout_type == LayoutType.step_workflow:
                _render_workflow(slide, spec_slide, theme)
            elif spec_slide.layout_type == LayoutType.architecture_layers:
                _render_layers(slide, spec_slide, theme)

        _add_text(slide, "DYNAMIC PRESENTATION ENGINE", 0.55, 7.1, 4.0, 0.2, size=8, color=theme["muted"], bold=True)
        _add_text(slide, str(spec_slide.slide_number), 12.15, 7.05, 0.55, 0.25, size=9, color=theme["muted"], align=PP_ALIGN.RIGHT)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()